from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.survey import SurveyDomain
from app.domain.survey_response import SurveyResponseDetailDomain
from app.services.exporters.base import SurveyResponseExporter


class XlsxSurveyResponseExporter(SurveyResponseExporter):
    """Exports responses as an .xlsx workbook: one sheet with response data
    (one row per response, one column per question), one sheet with a native
    Excel chart summarizing indicators."""

    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    file_extension = "xlsx"

    def export(
        self,
        survey: SurveyDomain,
        responses: list[SurveyResponseDetailDomain],
    ) -> bytes:
        workbook = Workbook()
        self._write_data_sheet(workbook.active, survey, responses)
        self._write_summary_sheet(workbook.create_sheet("Resumen"), responses)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _write_data_sheet(
        self,
        sheet: Worksheet,
        survey: SurveyDomain,
        responses: list[SurveyResponseDetailDomain],
    ) -> None:
        sheet.title = "Respuestas"
        headers = ["Completado por", "Fecha de relevamiento", "Score FIES"]
        headers += [q.text for q in survey.questions]
        sheet.append(headers)

        for response in responses:
            answer_by_question_id = {a.question_id: a for a in response.answers}
            row = [
                response.submitted_by.name,
                response.event_at.replace(tzinfo=None),
                response.score_fies if response.score_fies is not None else "",
            ]
            for question in survey.questions:
                answer = answer_by_question_id.get(question.id)
                if answer is None:
                    row.append("")
                    continue
                value = answer.value_label if answer.value_label else answer.value
                row.append(", ".join(value) if isinstance(value, list) else value)
            sheet.append(row)

    def _write_summary_sheet(
        self, sheet: Worksheet, responses: list[SurveyResponseDetailDomain]
    ) -> None:
        sheet.append(["Indicador", "Valor"])
        sheet.append(["Cantidad de respuestas", len(responses)])
        last_event_at = max((r.event_at for r in responses), default=None)
        sheet.append(
            [
                "Último relevamiento",
                last_event_at.replace(tzinfo=None) if last_event_at else "N/A",
            ]
        )

        fies_scores = [r.score_fies for r in responses if r.score_fies is not None]
        if fies_scores:
            self._write_fies_distribution(sheet, fies_scores)

    def _write_fies_distribution(
        self, sheet: Worksheet, fies_scores: list[int]
    ) -> None:
        start_row = sheet.max_row + 2
        sheet.cell(row=start_row, column=1, value="Score FIES")
        sheet.cell(row=start_row, column=2, value="Cantidad")
        for offset, score in enumerate(range(0, 9)):
            sheet.cell(row=start_row + 1 + offset, column=1, value=score)
            sheet.cell(
                row=start_row + 1 + offset, column=2, value=fies_scores.count(score)
            )

        chart = BarChart()
        chart.title = "Distribución de Score FIES"
        chart.y_axis.title = "Cantidad de respuestas"
        chart.x_axis.title = "Score FIES (0-8)"
        data = Reference(sheet, min_col=2, min_row=start_row, max_row=start_row + 9)
        categories = Reference(
            sheet, min_col=1, min_row=start_row + 1, max_row=start_row + 9
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        sheet.add_chart(chart, f"D{start_row}")
