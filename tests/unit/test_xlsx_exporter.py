import uuid
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.domain.common_domains import IdentifierRefDomain
from app.domain.survey import (
    QuestionDomain,
    ResponseType,
    SurveyDomain,
    SurveyInstrumentType,
    SurveyScope,
)
from app.domain.survey_response import (
    EstablishmentRefDomain,
    QuestionAnswerDomain,
    SurveyResponseDetailDomain,
)
from app.services.exporters.xlsx_exporter import XlsxSurveyResponseExporter

NOW = datetime(2026, 1, 1, tzinfo=timezone.utc)


def _survey_domain(fies: bool = False) -> SurveyDomain:
    question = QuestionDomain(
        id=uuid.uuid4(),
        text="¿Comió hoy?",
        response_type=ResponseType.SINGLE_CHOICE,
        is_required=True,
        position=0,
        options=[],
    )
    return SurveyDomain(
        id=uuid.uuid4(),
        name="Encuesta Test",
        description=None,
        scope=SurveyScope.GENERIC,
        instrument_type=SurveyInstrumentType.FIES
        if fies
        else SurveyInstrumentType.CUSTOM,
        created_by_id=uuid.uuid4(),
        is_active=True,
        response_count=0,
        questions=[question],
        establishment_ids=[],
        created_at=NOW,
        updated_at=NOW,
    )


def _detail_response(
    survey: SurveyDomain, score_fies: int | None = None
) -> SurveyResponseDetailDomain:
    question = survey.questions[0]
    return SurveyResponseDetailDomain(
        id=uuid.uuid4(),
        event_at=NOW,
        survey=IdentifierRefDomain(id=survey.id, name=survey.name),
        establishment=EstablishmentRefDomain(
            id=uuid.uuid4(),
            name="Comedor Test",
            city="Buenos Aires",
            province="Buenos Aires",
        ),
        submitted_by=IdentifierRefDomain(id=uuid.uuid4(), name="Juana Pérez"),
        score_fies=score_fies,
        answers=[
            QuestionAnswerDomain(
                question_id=question.id,
                question_text=question.text,
                response_type=question.response_type,
                value="opt-id",
                value_label="Sí",
            )
        ],
    )


class TestXlsxSurveyResponseExporter:
    def test_content_type_and_extension(self):
        exporter = XlsxSurveyResponseExporter()
        assert exporter.file_extension == "xlsx"
        assert "spreadsheetml" in exporter.content_type

    def test_export_produces_readable_workbook(self):
        survey = _survey_domain()
        responses = [_detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        file_bytes = exporter.export(survey, responses)

        workbook = load_workbook(BytesIO(file_bytes))
        assert "Respuestas" in workbook.sheetnames
        assert "Resumen" in workbook.sheetnames

    def test_data_sheet_has_one_row_per_response(self):
        survey = _survey_domain()
        responses = [_detail_response(survey), _detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Respuestas"]
        assert sheet.max_row == 3  # header + 2 responses

    def test_data_sheet_resolves_option_label(self):
        survey = _survey_domain()
        responses = [_detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Respuestas"]
        header = [c.value for c in sheet[1]]
        question_col = header.index(survey.questions[0].text) + 1
        assert sheet.cell(row=2, column=question_col).value == "Sí"

    def test_summary_sheet_has_response_count(self):
        survey = _survey_domain()
        responses = [_detail_response(survey), _detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Cantidad de respuestas", 2) in rows

    def test_fies_survey_includes_score_distribution(self):
        survey = _survey_domain(fies=True)
        responses = [
            _detail_response(survey, score_fies=3),
            _detail_response(survey, score_fies=3),
            _detail_response(survey, score_fies=7),
        ]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Score FIES", "Cantidad") in rows

    def test_custom_survey_has_no_score_distribution(self):
        survey = _survey_domain(fies=False)
        responses = [_detail_response(survey, score_fies=None)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Score FIES", "Cantidad") not in rows

    def test_empty_responses_produces_valid_workbook(self):
        survey = _survey_domain()
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, [])))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Cantidad de respuestas", 0) in rows
        assert ("Último relevamiento", "N/A") in rows
