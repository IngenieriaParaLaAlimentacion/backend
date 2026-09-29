import uuid
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.domain.common_domains import IdentifierRefDomain
from app.domain.subject import Sex
from app.domain.survey import (
    QuestionDomain,
    ResponseType,
    SurveyDomain,
    SurveyInstrumentType,
    SurveyResponseTarget,
)
from app.domain.survey_response import (
    EstablishmentRefDomain,
    QuestionAnswerDomain,
    SubjectRefDomain,
    SubmittedByRefDomain,
    SurveyResponseDetailDomain,
)
from app.services.exporters.xlsx_exporter import XlsxSurveyResponseExporter

NOW = datetime(2026, 1, 1, tzinfo=timezone.utc)


def _survey_domain(
    fies: bool = False,
    response_target: SurveyResponseTarget = SurveyResponseTarget.ANONYMOUS,
    requires_response_number: bool | None = True,
) -> SurveyDomain:
    question = QuestionDomain(
        id=uuid.uuid4(),
        text="¿Comió hoy?",
        response_type=ResponseType.SINGLE_CHOICE,
        is_required=True,
        position=0,
        options=[],
    )
    return SurveyDomain(
        id=uuid.uuid4(),
        name="Encuesta Test",
        description=None,
        response_target=response_target,
        instrument_type=SurveyInstrumentType.FIES
        if fies
        else SurveyInstrumentType.CUSTOM,
        requires_response_number=requires_response_number,
        created_by_id=uuid.uuid4(),
        is_active=True,
        response_count=0,
        questions=[question],
        created_at=NOW,
        updated_at=NOW,
    )


def _subject_ref() -> SubjectRefDomain:
    return SubjectRefDomain(
        id=uuid.uuid4(),
        dni=30111222,
        sex=Sex.FEMALE,
        name="María",
        last_name="Gómez",
        birth_date=date(1990, 5, 2),
    )


def _detail_response(
    survey: SurveyDomain,
    response_number: int = 1,
    score_fies: int | None = None,
    establishment: EstablishmentRefDomain | None = None,
    subject: SubjectRefDomain | None = None,
) -> SurveyResponseDetailDomain:
    question = survey.questions[0]
    if establishment is None and subject is None:
        establishment = EstablishmentRefDomain(
            id=uuid.uuid4(),
            name="Comedor Test",
            city="Buenos Aires",
            province="Buenos Aires",
        )
    return SurveyResponseDetailDomain(
        id=uuid.uuid4(),
        response_number=response_number,
        event_at=NOW,
        survey=IdentifierRefDomain(id=survey.id, name=survey.name),
        establishment=establishment,
        subject=subject,
        submitted_by=SubmittedByRefDomain(
            id=uuid.uuid4(), name="Juana", last_name="Pérez"
        ),
        score_fies=score_fies,
        answers=[
            QuestionAnswerDomain(
                question_id=question.id,
                question_text=question.text,
                response_type=question.response_type,
                value="opt-id",
                value_label="Sí",
            )
        ],
    )


class TestXlsxSurveyResponseExporter:
    def test_content_type_and_extension(self):
        exporter = XlsxSurveyResponseExporter()
        assert exporter.file_extension == "xlsx"
        assert "spreadsheetml" in exporter.content_type

    def test_data_sheet_has_one_row_per_response(self):
        survey = _survey_domain()
        responses = [_detail_response(survey), _detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Respuestas"]
        assert sheet.max_row == 3  # header + 2 responses

    @pytest.mark.parametrize("setting", [False, None])
    def test_optional_survey_omits_response_number_column(self, setting):
        survey = _survey_domain(requires_response_number=setting)
        responses = [_detail_response(survey, response_number=88)]

        workbook = load_workbook(
            BytesIO(XlsxSurveyResponseExporter().export(survey, responses))
        )
        sheet = workbook["Respuestas"]
        header = [cell.value for cell in sheet[1]]
        row = [cell.value for cell in sheet[2]]

        assert "N° de respuesta" not in header
        assert 88 not in row

    def test_data_sheet_resolves_option_label(self):
        survey = _survey_domain()
        responses = [_detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Respuestas"]
        header = [c.value for c in sheet[1]]
        question_col = header.index(survey.questions[0].text) + 1
        assert sheet.cell(row=2, column=question_col).value == "Sí"

    def test_summary_sheet_has_response_count(self):
        survey = _survey_domain()
        responses = [_detail_response(survey), _detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Cantidad de respuestas", 2) in rows

    def test_fies_survey_includes_score_distribution(self):
        survey = _survey_domain(fies=True)
        responses = [
            _detail_response(survey, score_fies=3),
            _detail_response(survey, score_fies=3),
            _detail_response(survey, score_fies=7),
        ]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Score FIES", "Cantidad") in rows

    def test_custom_survey_has_no_score_distribution(self):
        survey = _survey_domain(fies=False)
        responses = [_detail_response(survey, score_fies=None)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Score FIES", "Cantidad") not in rows

    def test_empty_responses_produces_valid_workbook(self):
        survey = _survey_domain()
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, [])))
        sheet = workbook["Resumen"]
        rows = list(sheet.iter_rows(values_only=True))
        assert ("Cantidad de respuestas", 0) in rows
        assert ("Último relevamiento", "N/A") in rows


class TestSubjectIdentityColumns:
    """A subject-targeted export identifies the person answering, so analysts
    can join it against that subject's measurement history."""

    def test_subject_targeted_export_has_identity_headers_first(self):
        survey = _survey_domain(response_target=SurveyResponseTarget.SUBJECT)
        responses = [_detail_response(survey, subject=_subject_ref())]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        header = [c.value for c in workbook["Respuestas"][1]]
        assert header[0] == "N° de respuesta"
        assert header[1:4] == ["Nombre", "Apellido", "DNI"]

    def test_subject_targeted_export_writes_identity_values(self):
        survey = _survey_domain(response_target=SurveyResponseTarget.SUBJECT)
        subject = _subject_ref()
        responses = [_detail_response(survey, subject=subject)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        row = [c.value for c in workbook["Respuestas"][2]]
        assert row[0] == responses[0].response_number
        assert row[1:4] == [subject.name, subject.last_name, subject.dni]

    @pytest.mark.parametrize(
        "target", [SurveyResponseTarget.ANONYMOUS, SurveyResponseTarget.ESTABLISHMENT]
    )
    def test_non_subject_export_has_no_identity_columns(self, target):
        survey = _survey_domain(response_target=target)
        responses = [_detail_response(survey)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        header = [c.value for c in workbook["Respuestas"][1]]
        assert header[0] == "N° de respuesta"
        assert header[1] == "Completado por"
        assert "DNI" not in header

    def test_missing_subject_writes_blanks_rather_than_raising(self):
        survey = _survey_domain(response_target=SurveyResponseTarget.SUBJECT)
        responses = [_detail_response(survey, subject=None)]
        exporter = XlsxSurveyResponseExporter()

        workbook = load_workbook(BytesIO(exporter.export(survey, responses)))
        row = [c.value for c in workbook["Respuestas"][2]]
        assert row[1:4] == [None, None, None]
