import uuid
from io import BytesIO

import pytest
import pytest_asyncio
from fastapi import status
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.domain.user import UserRole

pytestmark = pytest.mark.asyncio

SURVEYS_URL = "/api/v1/surveys"


def _responses_url(survey_id, *parts):
    segments = [SURVEYS_URL, str(survey_id), "responses"] + [str(p) for p in parts]
    return "/".join(segments)


NOW = "2024-06-01T12:00:00Z"


async def _make_establishment(place_id: str = "place-test-002"):
    from app.db.database import AsyncSessionLocal
    from app.models.establishments.establishment import Establishment

    async with AsyncSessionLocal() as session:
        establishment = Establishment(
            name="Otro Centro",
            formatted_address="Otra Calle 456",
            latitude=-31.4,
            longitude=-64.2,
            place_id=place_id,
            city="Cordoba",
            province="Cordoba",
        )
        session.add(establishment)
        await session.commit()
        await session.refresh(establishment)
        return establishment


async def _make_subject(dni: int = 33444555):
    from datetime import date

    from app.db.database import AsyncSessionLocal
    from app.models.subjects.subject import Sex, Subject

    async with AsyncSessionLocal() as session:
        subject = Subject(
            dni=dni,
            sex=Sex.MALE,
            name="Otro",
            last_name="Sujeto",
            birth_date=date(1985, 3, 14),
        )
        session.add(subject)
        await session.commit()
        await session.refresh(subject)
        return subject


@pytest_asyncio.fixture
async def admin_user(make_user_with_role):
    return await make_user_with_role(UserRole.ADMIN)


@pytest_asyncio.fixture
async def superadmin_user(make_user_with_role):
    return await make_user_with_role(UserRole.SUPERADMIN)


@pytest_asyncio.fixture
async def analyst_user(make_user_with_role):
    return await make_user_with_role(UserRole.ANALYST)


@pytest_asyncio.fixture
async def field_worker_without_access(make_user_with_role):
    return await make_user_with_role(UserRole.FIELD_WORKER)


async def _create_survey(client, headers, **overrides):
    body = {
        "name": overrides.pop("name", "Encuesta A"),
        "description": overrides.pop("description", None),
        "response_target": overrides.pop("response_target", "anonymous"),
        "requires_response_number": overrides.pop("requires_response_number", True),
        "questions": overrides.pop(
            "questions",
            [
                {
                    "text": "¿Comió hoy?",
                    "response_type": "single_choice",
                    "is_required": True,
                    "options": [{"text": "Sí"}, {"text": "No"}],
                },
                {
                    "text": "Comentarios",
                    "response_type": "text",
                    "is_required": False,
                    "options": [],
                },
            ],
        ),
    }
    response = await client.post(SURVEYS_URL, json=body, headers=headers)
    assert response.status_code == status.HTTP_201_CREATED
    return response.json()


def _valid_answers(survey: dict) -> dict:
    """A minimal valid answer set: only the required single_choice question."""
    question = survey["questions"][0]
    return {question["id"]: question["options"][0]["id"]}


def _response_body(survey_id, answers, response_number: int = 1, **overrides) -> dict:
    """Build a submission body for the default (anonymous) survey.

    Anonymous surveys carry no target, so nothing else is sent. Tests covering
    subject- or establishment-targeted surveys pass ``subject_id=`` or
    ``establishment_id=`` explicitly.
    """
    body = {
        "survey_id": str(survey_id),
        "answers": answers,
        "response_number": response_number,
        "event_at": NOW,
    }
    body.update(overrides)
    return body


class TestSubmitAuth:
    async def test_unauthenticated_returns_401(self, client, db_establishment):
        response = await client.post(
            _responses_url(uuid.uuid4()),
            json=_response_body(uuid.uuid4(), {}),
        )
        assert response.status_code == status.HTTP_401_UNAUTHORIZED

    async def test_analyst_forbidden(self, client, db_establishment, analyst_user):
        headers, _ = analyst_user
        response = await client.post(
            _responses_url(uuid.uuid4()),
            json=_response_body(uuid.uuid4(), {}),
            headers=headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN


class TestSubmitRoleMatrix:
    async def test_field_worker_without_any_establishment_succeeds(
        self, client, db_establishment, admin_user, field_worker_without_access
    ):
        """Previously 403. The establishment on a response is data, not
        permission, so a field worker who signed up purely to answer surveys —
        with no establishment relation at all — can submit."""
        admin_headers, _ = admin_user
        fw_headers, fw_id = field_worker_without_access
        survey = await _create_survey(client, admin_headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_201_CREATED
        assert response.json()["submitted_by_id"] == str(fw_id)

    async def test_field_worker_may_answer_for_an_unrelated_establishment(
        self, client, db_establishment, admin_user, field_worker_without_access
    ):
        """No approved relation is required to name an establishment."""
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_without_access
        survey = await _create_survey(
            client, admin_headers, response_target="establishment"
        )
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                establishment_id=str(db_establishment.id),
            ),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_201_CREATED
        assert response.json()["establishment_id"] == str(db_establishment.id)

    async def test_admin_succeeds(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )
        assert response.status_code == status.HTTP_201_CREATED

    async def test_superadmin_succeeds(
        self, client, db_establishment, admin_user, superadmin_user
    ):
        admin_headers, _ = admin_user
        sa_headers, _ = superadmin_user
        survey = await _create_survey(client, admin_headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=sa_headers,
        )
        assert response.status_code == status.HTTP_201_CREATED


class TestSubmitValidation:
    async def test_missing_required_question_returns_400(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], {}),
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_unknown_survey_returns_404(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        response = await client.post(
            _responses_url(uuid.uuid4()),
            json=_response_body(uuid.uuid4(), {}),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_unknown_establishment_returns_404(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="establishment")
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                establishment_id=str(uuid.uuid4()),
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_unknown_subject_returns_404(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                subject_id=str(uuid.uuid4()),
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    @pytest.mark.parametrize(
        "target,field",
        [
            ("subject", "establishment_id"),
            ("establishment", "subject_id"),
            ("anonymous", "subject_id"),
            ("anonymous", "establishment_id"),
        ],
    )
    async def test_wrong_entity_for_target_returns_400(
        self, client, db_establishment, admin_user, target, field
    ):
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta {target}"
        )
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"], _valid_answers(survey), **{field: str(uuid.uuid4())}
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    @pytest.mark.parametrize("target", ["subject", "establishment"])
    async def test_missing_required_entity_returns_400(
        self, client, admin_user, target
    ):
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta {target}"
        )
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_multiple_choice_with_non_string_element_returns_422(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(
            client,
            headers,
            questions=[
                {
                    "text": "¿Qué alimentos consumió?",
                    "response_type": "multiple_choice",
                    "is_required": True,
                    "options": [{"text": "Verduras"}, {"text": "Frutas"}],
                }
            ],
        )
        question = survey["questions"][0]
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"], {question["id"]: [{"not": "a valid option id"}]}
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT

    async def test_future_event_at_returns_422(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"], _valid_answers(survey), event_at="2099-01-01T00:00:00Z"
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT

    async def test_subject_targeted_submission_succeeds(
        self, client, admin_user, db_subject
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                subject_id=str(db_subject.id),
            ),
            headers=headers,
        )
        assert response.status_code == status.HTTP_201_CREATED
        body = response.json()
        assert body["subject_id"] == str(db_subject.id)
        assert body["establishment_id"] is None

    async def test_anonymous_submission_carries_no_target(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )
        assert response.status_code == status.HTTP_201_CREATED
        body = response.json()
        assert body["subject_id"] is None
        assert body["establishment_id"] is None


class TestTargetUniqueness:
    """A subject, and an establishment, answer a given survey once. Anonymous
    surveys stay unbounded."""

    async def test_second_response_for_same_subject_returns_409(
        self, client, admin_user, db_subject
    ):
        headers, user_id = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        body = _response_body(
            survey["id"], _valid_answers(survey), subject_id=str(db_subject.id)
        )
        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert first.status_code == status.HTTP_201_CREATED

        second = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert second.status_code == status.HTTP_409_CONFLICT
        assert second.headers["content-type"].startswith("application/problem+json")
        existing = second.json()["existing_response"]
        assert existing["id"] == first.json()["id"]
        assert existing["submitted_by"]["id"] == str(user_id)
        assert existing["event_at"] is not None

    async def test_second_response_for_same_establishment_returns_409(
        self, client, admin_user, db_establishment
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="establishment")
        body = _response_body(
            survey["id"],
            _valid_answers(survey),
            establishment_id=str(db_establishment.id),
        )
        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert first.status_code == status.HTTP_201_CREATED

        second = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert second.status_code == status.HTTP_409_CONFLICT

    async def test_anonymous_survey_accepts_many_responses(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        for _ in range(3):
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED

    async def test_required_survey_rejects_missing_number(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, requires_response_number=True)
        body = _response_body(survey["id"], _valid_answers(survey))
        del body["response_number"]

        response = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )

        assert response.status_code == status.HTTP_400_BAD_REQUEST
        assert response.json()["type"].endswith("/errors/invalid-survey-response")

    @pytest.mark.parametrize("setting", [False, None])
    async def test_optional_survey_persists_null_number(
        self, client, admin_user, setting
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, requires_response_number=setting)
        body = _response_body(
            survey["id"], _valid_answers(survey), response_number=None
        )

        response = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )

        assert response.status_code == status.HTTP_201_CREATED
        assert response.json()["response_number"] is None

    async def test_anonymous_responses_may_share_a_response_number(
        self, client, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        body = _response_body(survey["id"], _valid_answers(survey), response_number=1)

        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        second = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )

        assert first.status_code == status.HTTP_201_CREATED
        assert second.status_code == status.HTTP_201_CREATED
        assert first.json()["response_number"] == second.json()["response_number"] == 1

    @pytest.mark.parametrize("target", ["subject", "establishment"])
    async def test_targeted_responses_may_share_a_number_when_targets_differ(
        self, client, admin_user, db_establishment, db_subject, target
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target=target)
        if target == "subject":
            other_target = await _make_subject(dni=33444556)
            target_key = "subject_id"
            target_ids = (db_subject.id, other_target.id)
        else:
            other_target = await _make_establishment(place_id="place-test-003")
            target_key = "establishment_id"
            target_ids = (db_establishment.id, other_target.id)

        for target_id in target_ids:
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    response_number=1,
                    **{target_key: str(target_id)},
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED

    async def test_same_subject_may_answer_a_different_survey(
        self, client, admin_user, db_subject
    ):
        headers, _ = admin_user
        first_survey = await _create_survey(
            client, headers, name="Encuesta 1", response_target="subject"
        )
        second_survey = await _create_survey(
            client, headers, name="Encuesta 2", response_target="subject"
        )
        for survey in (first_survey, second_survey):
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    subject_id=str(db_subject.id),
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED

    async def test_replayed_id_wins_over_uniqueness(
        self, client, admin_user, db_subject
    ):
        """Id-first ordering: a retried sync gets its own response back with a
        200, not a conflict blaming it for its own earlier submission."""
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        client_id = str(uuid.uuid4())
        body = _response_body(
            survey["id"],
            _valid_answers(survey),
            subject_id=str(db_subject.id),
            id=client_id,
        )
        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert first.status_code == status.HTTP_201_CREATED

        replay = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert replay.status_code == status.HTTP_200_OK
        assert replay.json()["id"] == client_id

    async def test_conflict_survives_without_the_service_pre_check(
        self, client, admin_user, db_subject, monkeypatch
    ):
        """The partial unique indexes, not the pre-check, are the real
        guarantee. With the pre-check blinded — as it is in a genuine race —
        the repository must still recognise which index Postgres rejected the
        insert on, or the collision is mislabelled as an id replay and escapes
        as a 500 instead of this 409.
        """
        from app.repositories.surveys.survey_response_repository import (
            SurveyResponseRepository,
        )

        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        body = _response_body(
            survey["id"], _valid_answers(survey), subject_id=str(db_subject.id)
        )
        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert first.status_code == status.HTTP_201_CREATED

        real_get_by_target = SurveyResponseRepository.get_by_target
        calls = {"n": 0}

        async def blind_first_call(self, *args, **kwargs):
            """Return None once — the pre-check's view of a lost race — then
            behave normally so the post-INSERT lookup can name the winner."""
            calls["n"] += 1
            if calls["n"] == 1:
                return None
            return await real_get_by_target(self, *args, **kwargs)

        monkeypatch.setattr(SurveyResponseRepository, "get_by_target", blind_first_call)

        second = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert second.status_code == status.HTTP_409_CONFLICT
        assert second.json()["existing_response"]["id"] == first.json()["id"]

    async def test_id_replayed_under_a_different_survey_returns_400(
        self, client, admin_user
    ):
        """One client id belongs to one survey. Reusing it elsewhere would hand
        back the first survey's answers as a 200."""
        headers, _ = admin_user
        first_survey = await _create_survey(client, headers, name="Encuesta 1")
        second_survey = await _create_survey(client, headers, name="Encuesta 2")
        client_id = str(uuid.uuid4())

        created = await client.post(
            _responses_url(first_survey["id"]),
            json=_response_body(
                first_survey["id"], _valid_answers(first_survey), id=client_id
            ),
            headers=headers,
        )
        assert created.status_code == status.HTTP_201_CREATED

        crossed = await client.post(
            _responses_url(second_survey["id"]),
            json=_response_body(
                second_survey["id"], _valid_answers(second_survey), id=client_id
            ),
            headers=headers,
        )
        assert crossed.status_code == status.HTTP_400_BAD_REQUEST


class TestSubmitIdempotentReplay:
    async def test_same_client_id_replays_without_duplicate_row(
        self, client, db_establishment, admin_user
    ):
        from app.models.surveys.survey_response import SurveyResponse

        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        response_id = str(uuid.uuid4())
        body = _response_body(survey["id"], _valid_answers(survey), id=response_id)

        first = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert first.status_code == status.HTTP_201_CREATED

        second = await client.post(
            _responses_url(survey["id"]), json=body, headers=headers
        )
        assert second.status_code == status.HTTP_200_OK
        assert second.json() == first.json()

        from app.db.database import AsyncSessionLocal

        async with AsyncSessionLocal() as session:
            count = await session.execute(
                select(func.count()).select_from(SurveyResponse)
            )
            assert count.scalar_one() == 1


class TestResponseCount:
    async def test_submit_increments_survey_response_count(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )
        get_response = await client.get(
            f"{SURVEYS_URL}/{survey['id']}", headers=headers
        )
        assert get_response.json()["response_count"] == 1


class TestEdit:
    async def test_optional_edit_clears_live_and_audit_number(
        self, client, db_establishment, admin_user
    ):
        from app.db.database import AsyncSessionLocal
        from app.models.surveys.survey_response import SurveyResponse
        from app.models.surveys.survey_response_edit import SurveyResponseEdit

        headers, _ = admin_user
        survey = await _create_survey(client, headers, requires_response_number=False)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"], _valid_answers(survey), response_number=12
                ),
                headers=headers,
            )
        ).json()

        edited = await client.put(
            _responses_url(survey["id"], created["id"]),
            json={"answers": _valid_answers(survey), "event_at": NOW},
            headers=headers,
        )

        assert edited.status_code == status.HTTP_200_OK
        assert edited.json()["response_number"] is None
        async with AsyncSessionLocal() as session:
            live = await session.get(SurveyResponse, uuid.UUID(created["id"]))
            audit = await session.scalar(
                select(SurveyResponseEdit).where(
                    SurveyResponseEdit.response_id == uuid.UUID(created["id"])
                )
            )
        assert live.response_number is None
        assert audit.response_number is None

    async def test_required_edit_without_number_creates_no_audit(
        self, client, db_establishment, admin_user
    ):
        from app.db.database import AsyncSessionLocal
        from app.models.surveys.survey_response_edit import SurveyResponseEdit

        headers, _ = admin_user
        survey = await _create_survey(client, headers, requires_response_number=True)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=headers,
            )
        ).json()

        edited = await client.put(
            _responses_url(survey["id"], created["id"]),
            json={"answers": _valid_answers(survey), "event_at": NOW},
            headers=headers,
        )

        assert edited.status_code == status.HTTP_400_BAD_REQUEST
        async with AsyncSessionLocal() as session:
            count = await session.scalar(
                select(func.count()).select_from(SurveyResponseEdit)
            )
        assert count == 0

    async def test_missing_response_returns_404(self, client, admin_user):
        headers, _ = admin_user
        response = await client.put(
            _responses_url(uuid.uuid4(), uuid.uuid4()),
            json={"answers": {}, "response_number": 1, "event_at": NOW},
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_non_owner_field_worker_forbidden(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=admin_headers,
            )
        ).json()
        response = await client.put(
            _responses_url(created["survey_id"], created["id"]),
            json={
                "answers": _valid_answers(survey),
                "response_number": created["response_number"],
                "event_at": NOW,
            },
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_owner_field_worker_can_edit(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=fw_headers,
            )
        ).json()
        question = survey["questions"][0]
        new_answers = {question["id"]: question["options"][1]["id"]}
        response = await client.put(
            _responses_url(created["survey_id"], created["id"]),
            json={
                "answers": new_answers,
                "response_number": created["response_number"],
                "event_at": NOW,
            },
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_200_OK
        assert response.json()["answers"] == new_answers

    async def test_admin_can_override_any_response(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=fw_headers,
            )
        ).json()
        response = await client.put(
            _responses_url(created["survey_id"], created["id"]),
            json={
                "answers": _valid_answers(survey),
                "response_number": created["response_number"],
                "event_at": NOW,
            },
            headers=admin_headers,
        )
        assert response.status_code == status.HTTP_200_OK

    async def test_edit_appends_audit_row(self, client, db_establishment, admin_user):
        from app.models.surveys.survey_response_edit import SurveyResponseEdit

        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=headers,
            )
        ).json()
        await client.put(
            _responses_url(created["survey_id"], created["id"]),
            json={
                "answers": _valid_answers(survey),
                "response_number": created["response_number"],
                "event_at": NOW,
            },
            headers=headers,
        )

        from app.db.database import AsyncSessionLocal

        async with AsyncSessionLocal() as session:
            count = await session.execute(
                select(func.count()).select_from(SurveyResponseEdit)
            )
            assert count.scalar_one() == 1

    async def test_edit_to_a_used_number_succeeds_and_preserves_audit_snapshot(
        self, client, db_establishment, admin_user
    ):
        from app.db.database import AsyncSessionLocal
        from app.models.surveys.survey_response_edit import SurveyResponseEdit

        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        answers = _valid_answers(survey)
        first = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], answers, response_number=1),
            headers=headers,
        )
        second = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], answers, response_number=2),
            headers=headers,
        )
        assert first.status_code == status.HTTP_201_CREATED
        assert second.status_code == status.HTTP_201_CREATED

        edited = await client.put(
            _responses_url(survey["id"], second.json()["id"]),
            json={
                "answers": answers,
                "response_number": 1,
                "event_at": NOW,
            },
            headers=headers,
        )

        assert edited.status_code == status.HTTP_200_OK
        assert edited.json()["response_number"] == 1
        async with AsyncSessionLocal() as session:
            audit = await session.scalar(
                select(SurveyResponseEdit).where(
                    SurveyResponseEdit.response_id == uuid.UUID(second.json()["id"])
                )
            )
        assert audit is not None
        assert audit.response_number == 1

    async def test_last_write_wins_older_event_at_does_not_overwrite(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        original_answers = _valid_answers(survey)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"], original_answers, event_at="2024-06-02T00:00:00Z"
                ),
                headers=headers,
            )
        ).json()
        question = survey["questions"][0]
        older_answers = {question["id"]: question["options"][1]["id"]}

        response = await client.put(
            _responses_url(created["survey_id"], created["id"]),
            json={
                "answers": older_answers,
                "response_number": created["response_number"],
                "event_at": "2024-06-01T00:00:00Z",
            },
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK
        # The edit is logged, but the response's current state is unchanged
        # since its event_at is more recent than the edit's.
        assert response.json()["answers"] == original_answers

        from app.db.database import AsyncSessionLocal
        from app.models.surveys.survey_response_edit import SurveyResponseEdit

        async with AsyncSessionLocal() as session:
            count = await session.execute(
                select(func.count()).select_from(SurveyResponseEdit)
            )
            assert count.scalar_one() == 1


class TestGetResponse:
    """The full-detail view a mobile client hits after tapping a `/me/activity`
    entry, since the activity feed only carries a summary."""

    async def test_missing_response_returns_404(self, client, admin_user):
        headers, _ = admin_user
        response = await client.get(_responses_url(uuid.uuid4()), headers=headers)
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_owner_field_worker_can_view_own_response(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        answers = _valid_answers(survey)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], answers),
                headers=fw_headers,
            )
        ).json()
        response = await client.get(
            _responses_url(created["survey_id"], created["id"]), headers=fw_headers
        )
        assert response.status_code == status.HTTP_200_OK
        assert response.json()["answers"] == answers

    async def test_non_owner_field_worker_forbidden(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=admin_headers,
            )
        ).json()
        response = await client.get(
            _responses_url(created["survey_id"], created["id"]), headers=fw_headers
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_admin_can_view_any_response(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=fw_headers,
            )
        ).json()
        response = await client.get(
            _responses_url(created["survey_id"], created["id"]), headers=admin_headers
        )
        assert response.status_code == status.HTTP_200_OK

    async def test_analyst_forbidden(
        self, client, db_establishment, admin_user, analyst_user
    ):
        admin_headers, _ = admin_user
        analyst_headers, _ = analyst_user
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=admin_headers,
            )
        ).json()
        response = await client.get(
            _responses_url(created["survey_id"], created["id"]), headers=analyst_headers
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN


class TestListAvailable:
    """Availability stopped being establishment-scoped: every role sees every
    active survey, whatever its response target."""

    async def test_field_worker_with_access_sees_every_active_survey(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access

        await _create_survey(client, admin_headers, name="Anonymous Survey")
        await _create_survey(
            client,
            admin_headers,
            name="Subject Survey",
            response_target="subject",
        )
        await _create_survey(
            client,
            admin_headers,
            name="Establishment Survey",
            response_target="establishment",
        )

        response = await client.get(f"{SURVEYS_URL}/available", headers=fw_headers)
        assert response.status_code == status.HTTP_200_OK
        names = {s["name"] for s in response.json()["items"]}
        assert names == {
            "Anonymous Survey",
            "Subject Survey",
            "Establishment Survey",
        }

    async def test_field_worker_without_any_establishment_sees_the_same(
        self, client, db_establishment, admin_user, field_worker_without_access
    ):
        """Previously this worker saw generic surveys only."""
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_without_access
        await _create_survey(client, admin_headers, name="Anonymous Survey")
        await _create_survey(
            client,
            admin_headers,
            name="Establishment Survey",
            response_target="establishment",
        )
        response = await client.get(f"{SURVEYS_URL}/available", headers=fw_headers)
        names = {s["name"] for s in response.json()["items"]}
        assert names == {"Anonymous Survey", "Establishment Survey"}

    async def test_items_expose_response_target(
        self, client, admin_user, field_worker_without_access
    ):
        """Mobile routes on this before opening the form, so it must be on the
        list item and not only on the detail."""
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_without_access
        await _create_survey(
            client, admin_headers, name="Subject Survey", response_target="subject"
        )
        response = await client.get(f"{SURVEYS_URL}/available", headers=fw_headers)
        assert response.json()["items"][0]["response_target"] == "subject"

    async def test_inactive_surveys_excluded(
        self, client, admin_user, field_worker_without_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_without_access
        survey = await _create_survey(client, admin_headers, name="To Deactivate")
        await client.patch(
            f"{SURVEYS_URL}/{survey['id']}/status",
            json={"is_active": False},
            headers=admin_headers,
        )
        response = await client.get(f"{SURVEYS_URL}/available", headers=fw_headers)
        assert response.json()["total"] == 0

    async def test_analyst_allowed(self, client, analyst_user):
        headers, _ = analyst_user
        response = await client.get(f"{SURVEYS_URL}/available", headers=headers)
        assert response.status_code == status.HTTP_200_OK


class TestListBySurvey:
    async def test_shows_submitter_establishment_and_date(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, fw_id = field_worker_with_access
        survey = await _create_survey(
            client, admin_headers, response_target="establishment"
        )
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    establishment_id=str(db_establishment.id),
                ),
                headers=fw_headers,
            )
        ).json()

        response = await client.get(
            _responses_url(survey["id"]),
            headers=admin_headers,
        )

        assert response.status_code == status.HTTP_200_OK
        body = response.json()
        assert body["total"] == 1
        item = body["items"][0]
        assert item["id"] == created["id"]
        assert item["survey"]["id"] == survey["id"]
        assert item["establishment"]["id"] == str(db_establishment.id)
        assert item["subject"] is None
        assert item["submitted_by"]["id"] == str(fw_id)
        assert "event_at" in item

    async def test_orders_by_event_at_desc(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        older = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    event_at="2024-06-01T00:00:00Z",
                ),
                headers=headers,
            )
        ).json()
        newer = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    id=str(uuid.uuid4()),
                    event_at="2024-06-05T00:00:00Z",
                ),
                headers=headers,
            )
        ).json()

        response = await client.get(
            _responses_url(survey["id"]),
            headers=headers,
        )

        ids = [item["id"] for item in response.json()["items"]]
        assert ids == [newer["id"], older["id"]]

    @pytest.mark.parametrize(
        "sort_order,oldest_first",
        [
            ("event_at_asc", True),
            ("event_at_desc", False),
            ("created_at_asc", True),
            ("created_at_desc", False),
        ],
    )
    async def test_sort_order(
        self, client, db_establishment, admin_user, sort_order, oldest_first
    ):
        """Responses are created oldest event_at first, so creation order and
        event_at order agree — one expectation covers both sort columns."""
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        created = []
        for event_at in ("2024-06-01T00:00:00Z", "2024-06-05T00:00:00Z"):
            created.append(
                (
                    await client.post(
                        _responses_url(survey["id"]),
                        json=_response_body(
                            survey["id"],
                            _valid_answers(survey),
                            id=str(uuid.uuid4()),
                            event_at=event_at,
                        ),
                        headers=headers,
                    )
                ).json()
            )

        response = await client.get(
            _responses_url(survey["id"]),
            params={"sort_order": sort_order},
            headers=headers,
        )

        ids = [item["id"] for item in response.json()["items"]]
        expected = [row["id"] for row in created]
        assert ids == (expected if oldest_first else expected[::-1])

    async def test_rejects_unknown_sort_order(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)

        response = await client.get(
            _responses_url(survey["id"]),
            params={"sort_order": "subject_name_asc"},
            headers=headers,
        )

        assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT

    async def test_rejects_inverted_event_at_range(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)

        response = await client.get(
            _responses_url(survey["id"]),
            params={
                "event_at_from": "2024-06-05T00:00:00Z",
                "event_at_to": "2024-06-01T00:00:00Z",
            },
            headers=headers,
        )

        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_paginates(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        for _ in range(3):
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"], _valid_answers(survey), id=str(uuid.uuid4())
                ),
                headers=headers,
            )

        response = await client.get(
            _responses_url(survey["id"]),
            params={"page": 1, "page_size": 2},
            headers=headers,
        )

        body = response.json()
        assert len(body["items"]) == 2
        assert body["total"] == 3
        assert body["total_pages"] == 2

    async def test_only_returns_responses_for_requested_survey(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey_a = await _create_survey(client, headers, name="Encuesta A")
        survey_b = await _create_survey(client, headers, name="Encuesta B")
        await client.post(
            _responses_url(survey_a["id"]),
            json=_response_body(survey_a["id"], _valid_answers(survey_a)),
            headers=headers,
        )
        await client.post(
            _responses_url(survey_b["id"]),
            json=_response_body(survey_b["id"], _valid_answers(survey_b)),
            headers=headers,
        )

        response = await client.get(
            _responses_url(survey_a["id"]),
            headers=headers,
        )

        assert response.json()["total"] == 1

    async def test_field_worker_forbidden(
        self, client, admin_user, field_worker_without_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_without_access
        survey = await _create_survey(client, admin_headers)

        response = await client.get(
            _responses_url(survey["id"]),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_analyst_allowed(self, client, admin_user, analyst_user):
        admin_headers, _ = admin_user
        analyst_headers, _ = analyst_user
        survey = await _create_survey(client, admin_headers)

        response = await client.get(
            _responses_url(survey["id"]),
            headers=analyst_headers,
        )
        assert response.status_code == status.HTTP_200_OK

    async def test_unknown_survey_returns_404(self, client, admin_user):
        headers, _ = admin_user

        unknown_id = str(uuid.uuid4())
        response = await client.get(
            _responses_url(unknown_id),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_unauthenticated_returns_401(self, client, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)

        response = await client.get(_responses_url(survey["id"]))
        assert response.status_code == status.HTTP_401_UNAUTHORIZED


class TestGetDashboardDetail:
    async def test_shows_question_text_and_resolved_option_label(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        answers = _valid_answers(survey)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], answers),
                headers=headers,
            )
        ).json()

        response = await client.get(
            _responses_url(created["survey_id"], created["id"], "detail"),
            headers=headers,
        )

        assert response.status_code == status.HTTP_200_OK
        body = response.json()
        assert len(body["answers"]) == len(survey["questions"])
        first_question = survey["questions"][0]
        answer = next(
            a for a in body["answers"] if a["question_id"] == first_question["id"]
        )
        assert answer["question_text"] == first_question["text"]
        expected_option = next(
            o
            for o in first_question["options"]
            if o["id"] == answers[first_question["id"]]
        )
        assert answer["value_label"] == expected_option["text"]

    async def test_unanswered_optional_question_has_null_value(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=headers,
            )
        ).json()

        response = await client.get(
            _responses_url(created["survey_id"], created["id"], "detail"),
            headers=headers,
        )

        comments_question = survey["questions"][1]
        answer = next(
            a
            for a in response.json()["answers"]
            if a["question_id"] == comments_question["id"]
        )
        assert answer["value"] is None
        assert answer["value_label"] is None

    async def test_field_worker_forbidden(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], _valid_answers(survey)),
                headers=fw_headers,
            )
        ).json()

        response = await client.get(
            _responses_url(created["survey_id"], created["id"], "detail"),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_response_not_found_returns_404(self, client, admin_user):
        headers, _ = admin_user
        response = await client.get(
            _responses_url(uuid.uuid4()),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND


class TestFiesScoring:
    @staticmethod
    def _fies_survey_body(name: str) -> dict:
        from app.reference_data.fies_template import build_fies_template

        template = build_fies_template()
        return {
            "name": name,
            "description": template.description,
            "response_target": template.response_target.value,
            "instrument_type": "fies",
            "questions": [
                {
                    "text": q.text,
                    "response_type": q.response_type.value,
                    "is_required": q.is_required,
                    "options": [{"text": o.text} for o in q.options],
                }
                for q in template.questions
            ],
        }

    async def test_fies_response_gets_partial_score(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = (
            await client.post(
                SURVEYS_URL,
                json=self._fies_survey_body("FIES Julio 2026"),
                headers=headers,
            )
        ).json()

        answers = {
            q["id"]: q["options"][0]["id"] if i < 3 else q["options"][1]["id"]
            for i, q in enumerate(survey["questions"])
        }
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], answers),
            headers=headers,
        )
        assert response.status_code == status.HTTP_201_CREATED
        assert response.json()["score_fies"] == 3

    async def test_fies_response_all_affirmative_scores_eight(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = (
            await client.post(
                SURVEYS_URL,
                json=self._fies_survey_body("FIES All Yes"),
                headers=headers,
            )
        ).json()
        answers = {q["id"]: q["options"][0]["id"] for q in survey["questions"]}

        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], answers),
            headers=headers,
        )
        assert response.json()["score_fies"] == 8

    async def test_custom_survey_response_has_null_score(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)  # default instrument_type=custom
        response = await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )
        assert response.json()["score_fies"] is None

    async def test_editing_answers_recalculates_score(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = (
            await client.post(
                SURVEYS_URL,
                json=self._fies_survey_body("FIES Edit Test"),
                headers=headers,
            )
        ).json()
        original_answers = {
            q["id"]: q["options"][1]["id"] for q in survey["questions"]
        }  # all "No"
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], original_answers),
                headers=headers,
            )
        ).json()
        assert created["score_fies"] == 0

        new_answers = {
            q["id"]: q["options"][0]["id"] for q in survey["questions"]
        }  # all "Sí"
        edited = (
            await client.put(
                _responses_url(created["survey_id"], created["id"]),
                json={
                    "answers": new_answers,
                    "response_number": created["response_number"],
                    "event_at": NOW,
                },
                headers=headers,
            )
        ).json()
        assert edited["score_fies"] == 8

    async def test_score_appears_in_dashboard_list(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = (
            await client.post(
                SURVEYS_URL,
                json=self._fies_survey_body("FIES Dashboard List"),
                headers=headers,
            )
        ).json()
        answers = {
            q["id"]: q["options"][1]["id"] for q in survey["questions"]
        }  # all "No"
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], answers),
            headers=headers,
        )

        list_response = await client.get(_responses_url(survey["id"]), headers=headers)
        assert list_response.json()["items"][0]["score_fies"] == 0

    async def test_score_appears_in_response_detail(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = (
            await client.post(
                SURVEYS_URL,
                json=self._fies_survey_body("FIES Detail View"),
                headers=headers,
            )
        ).json()
        answers = {
            q["id"]: q["options"][0]["id"] if i < 5 else q["options"][1]["id"]
            for i, q in enumerate(survey["questions"])
        }
        created = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(survey["id"], answers),
                headers=headers,
            )
        ).json()

        detail = await client.get(
            _responses_url(created["survey_id"], created["id"]),
            headers=headers,
        )
        assert detail.json()["score_fies"] == 5


class TestListFilters:
    """Geographic and establishment filters only make sense for
    establishment-targeted surveys, so these all build one."""

    @staticmethod
    async def _survey_with_two_establishments(client, headers, db_establishment, other):
        survey = await _create_survey(client, headers, response_target="establishment")
        for establishment_id in (db_establishment.id, other.id):
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    establishment_id=str(establishment_id),
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED
        return survey

    async def test_filters_by_establishment(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        other = await _make_establishment(place_id="place-filter-001")
        survey = await self._survey_with_two_establishments(
            client, headers, db_establishment, other
        )

        response = await client.get(
            _responses_url(survey["id"]),
            params={
                "establishment_id": str(db_establishment.id),
                "page": 1,
                "page_size": 10,
            },
            headers=headers,
        )
        body = response.json()
        assert body["total"] == 1
        assert body["items"][0]["establishment"]["id"] == str(db_establishment.id)

    async def test_filters_by_subject(self, client, admin_user, db_subject):
        headers, _ = admin_user
        other = await _make_subject()
        survey = await _create_survey(client, headers, response_target="subject")
        for subject_id in (db_subject.id, other.id):
            created = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    id=str(uuid.uuid4()),
                    subject_id=str(subject_id),
                ),
                headers=headers,
            )
            assert created.status_code == status.HTTP_201_CREATED

        response = await client.get(
            _responses_url(survey["id"]),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        body = response.json()
        assert body["total"] == 1
        assert body["items"][0]["subject"]["id"] == str(db_subject.id)

    @pytest.mark.parametrize("target", ["establishment", "anonymous"])
    async def test_subject_filter_on_non_subject_survey_returns_400(
        self, client, admin_user, db_subject, target
    ):
        """Mirror of the establishment case: ``subject_id`` is null on every
        response to a survey not targeted at subjects."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta sujeto {target}"
        )
        response = await client.get(
            _responses_url(survey["id"]),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    @pytest.mark.parametrize("target", ["subject", "anonymous"])
    @pytest.mark.parametrize("field", ["city", "province"])
    async def test_location_filter_on_non_establishment_survey_returns_400(
        self, client, admin_user, target, field
    ):
        """Those responses carry no establishment, so the join would silently
        drop them — an empty page reads as "no data", not "wrong filter"."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta {target}"
        )
        response = await client.get(
            _responses_url(survey["id"]),
            params={field: "Cordoba"},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_filters_by_event_at_range(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        old = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    event_at="2024-01-01T00:00:00Z",
                ),
                headers=headers,
            )
        ).json()
        recent = (
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    id=str(uuid.uuid4()),
                    event_at="2024-06-15T00:00:00Z",
                ),
                headers=headers,
            )
        ).json()

        response = await client.get(
            _responses_url(survey["id"]),
            params={
                "survey_id": survey["id"],
                "event_at_from": "2024-05-01T00:00:00Z",
            },
            headers=headers,
        )
        ids = [item["id"] for item in response.json()["items"]]
        assert ids == [recent["id"]]
        assert old["id"] not in ids

    async def test_response_number_exact_and_range_filters_include_duplicates(
        self, client, db_establishment, admin_user
    ):
        """Number filters return every response sharing the requested number."""
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        responses = []
        for response_number in (74, 74, 75):
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    response_number=response_number,
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED
            responses.append(response.json())

        exact = await client.get(
            _responses_url(survey["id"]),
            params={"response_number": 74},
            headers=headers,
        )
        assert exact.status_code == status.HTTP_200_OK
        assert {item["id"] for item in exact.json()["items"]} == {
            responses[0]["id"],
            responses[1]["id"],
        }

        ranged = await client.get(
            _responses_url(survey["id"]),
            params={"response_number_from": 74, "response_number_to": 75},
            headers=headers,
        )
        assert ranged.status_code == status.HTTP_200_OK
        assert {item["id"] for item in ranged.json()["items"]} == {
            item["id"] for item in responses
        }

    async def test_filters_by_city(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        cordoba = await _make_establishment(place_id="place-filter-002")
        survey = await self._survey_with_two_establishments(
            client, headers, db_establishment, cordoba
        )

        response = await client.get(
            _responses_url(survey["id"]),
            params={"city": "Buenos Aires"},
            headers=headers,
        )
        body = response.json()
        assert body["total"] == 1
        assert body["items"][0]["establishment"]["city"] == "Buenos Aires"

    async def test_filters_by_province(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        cordoba = await _make_establishment(place_id="place-filter-003")
        survey = await self._survey_with_two_establishments(
            client, headers, db_establishment, cordoba
        )

        response = await client.get(
            _responses_url(survey["id"]),
            params={"province": "Cordoba"},
            headers=headers,
        )
        body = response.json()
        assert body["total"] == 1
        assert body["items"][0]["establishment"]["province"] == "Cordoba"

    async def test_city_filter_case_insensitive(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="establishment")
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                establishment_id=str(db_establishment.id),
            ),
            headers=headers,
        )

        response = await client.get(
            _responses_url(survey["id"]),
            params={"city": "buenos aires"},  # lowercase
            headers=headers,
        )
        assert response.json()["total"] == 1

    async def test_no_filters_returns_all(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )

        response = await client.get(_responses_url(survey["id"]), headers=headers)
        assert response.json()["total"] == 1


ANALYTICS_URL = "/api/v1/analytics"


class TestSurveyResponseSummary:
    async def test_counts_all_responses_with_no_filters(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"], _valid_answers(survey), event_at="2024-01-01T00:00:00Z"
            ),
            headers=headers,
        )
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                id=str(uuid.uuid4()),
                event_at="2024-06-01T00:00:00Z",
            ),
            headers=headers,
        )

        url = f"{ANALYTICS_URL}/survey-response-summary"
        response = await client.get(url, headers=headers)

        assert response.status_code == status.HTTP_200_OK
        body = response.json()
        assert body["response_count"] == 2
        # Anonymous responses have no establishment, so the distinct count is
        # legitimately zero — the metric only means something for
        # establishment-targeted surveys.
        assert body["participating_establishments_count"] == 0
        assert body["last_event_at"] == "2024-06-01T00:00:00Z"

    async def test_counts_distinct_establishments(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        other = await _make_establishment(place_id="place-indicators-001")
        survey = await _create_survey(client, headers, response_target="establishment")

        for establishment_id in (db_establishment.id, other.id):
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    establishment_id=str(establishment_id),
                ),
                headers=headers,
            )

        url = f"{ANALYTICS_URL}/survey-response-summary"
        response = await client.get(url, headers=headers)

        assert response.json()["participating_establishments_count"] == 2

    async def test_filters_by_survey_id(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey_a = await _create_survey(client, headers, name="Encuesta A")
        survey_b = await _create_survey(client, headers, name="Encuesta B")
        await client.post(
            _responses_url(survey_a["id"]),
            json=_response_body(survey_a["id"], _valid_answers(survey_a)),
            headers=headers,
        )
        await client.post(
            _responses_url(survey_b["id"]),
            json=_response_body(survey_b["id"], _valid_answers(survey_b)),
            headers=headers,
        )

        response = await client.get(
            f"{ANALYTICS_URL}/survey-response-summary",
            params={"survey_id": survey_a["id"]},
            headers=headers,
        )

        assert response.json()["response_count"] == 1

    async def test_filters_by_event_at_range(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"], _valid_answers(survey), event_at="2024-01-01T00:00:00Z"
            ),
            headers=headers,
        )
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                id=str(uuid.uuid4()),
                event_at="2024-06-15T00:00:00Z",
            ),
            headers=headers,
        )

        response = await client.get(
            f"{ANALYTICS_URL}/survey-response-summary",
            params={"event_at_from": "2024-05-01T00:00:00Z"},
            headers=headers,
        )

        body = response.json()
        assert body["response_count"] == 1
        assert body["last_event_at"] == "2024-06-15T00:00:00Z"

    async def test_filters_by_city(self, client, db_establishment, admin_user):
        """The global summary keeps its establishment join, so a location
        filter there simply matches only establishment-linked responses."""
        headers, _ = admin_user
        cordoba = await _make_establishment(place_id="place-indicators-002")
        survey = await _create_survey(client, headers, response_target="establishment")
        for establishment_id in (db_establishment.id, cordoba.id):
            await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    establishment_id=str(establishment_id),
                ),
                headers=headers,
            )

        response = await client.get(
            f"{ANALYTICS_URL}/survey-response-summary",
            params={"city": "Buenos Aires"},
            headers=headers,
        )

        assert response.json()["response_count"] == 1

    async def test_no_responses_returns_zero_and_null_last_event_at(
        self, client, admin_user
    ):
        headers, _ = admin_user
        url = f"{ANALYTICS_URL}/survey-response-summary"
        response = await client.get(url, headers=headers)

        body = response.json()
        assert body["response_count"] == 0
        assert body["participating_establishments_count"] == 0
        assert body["last_event_at"] is None

    async def test_invalid_date_range_returns_400(self, client, admin_user):
        headers, _ = admin_user
        response = await client.get(
            f"{ANALYTICS_URL}/survey-response-summary",
            params={
                "event_at_from": "2024-06-01T00:00:00Z",
                "event_at_to": "2024-01-01T00:00:00Z",
            },
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_field_worker_forbidden(self, client, field_worker_without_access):
        headers, _ = field_worker_without_access
        url = f"{ANALYTICS_URL}/survey-response-summary"
        response = await client.get(url, headers=headers)
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_analyst_allowed(self, client, analyst_user):
        headers, _ = analyst_user
        url = f"{ANALYTICS_URL}/survey-response-summary"
        response = await client.get(url, headers=headers)
        assert response.status_code == status.HTTP_200_OK

    async def test_unauthenticated_returns_401(self, client):
        response = await client.get(f"{ANALYTICS_URL}/survey-response-summary")
        assert response.status_code == status.HTTP_401_UNAUTHORIZED


class TestExport:
    async def test_admin_can_export_xlsx(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )

        response = await client.get(
            _responses_url(survey["id"], "export"),
            headers=headers,
        )

        assert response.status_code == status.HTTP_200_OK
        assert "spreadsheetml" in response.headers["content-type"]
        assert response.content[:2] == b"PK"  # xlsx is a zip archive

    async def test_response_number_range_export_includes_duplicates(
        self, client, db_establishment, admin_user
    ):
        """The number range export retains all matching duplicate labels."""
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        for response_number in (74, 74, 75):
            response = await client.post(
                _responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    response_number=response_number,
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED

        response = await client.get(
            _responses_url(survey["id"], "export"),
            params={"response_number_from": 74, "response_number_to": 74},
            headers=headers,
        )

        assert response.status_code == status.HTTP_200_OK
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook["Respuestas"].iter_rows(min_row=2, values_only=True))
        assert [row[0] for row in rows] == [74, 74]

    async def test_filters_by_establishment(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="establishment")
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                establishment_id=str(db_establishment.id),
            ),
            headers=headers,
        )

        response = await client.get(
            _responses_url(survey["id"], "export"),
            params={"establishment_id": str(db_establishment.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK

    async def test_filters_by_subject(self, client, admin_user, db_subject):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        await client.post(
            _responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                subject_id=str(db_subject.id),
            ),
            headers=headers,
        )

        response = await client.get(
            _responses_url(survey["id"], "export"),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK

    @pytest.mark.parametrize("target", ["establishment", "anonymous"])
    async def test_subject_filter_on_non_subject_survey_returns_400(
        self, client, admin_user, db_subject, target
    ):
        """Nothing in this survey carries a subject, so the filter would export
        an empty workbook that reads as "no responses"."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta sujeto {target}"
        )

        response = await client.get(
            _responses_url(survey["id"], "export"),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    @pytest.mark.parametrize("target", ["subject", "anonymous"])
    async def test_establishment_filter_on_non_establishment_survey_returns_400(
        self, client, db_establishment, admin_user, target
    ):
        """Nothing in this survey carries an establishment, so the filter would
        export an empty workbook that reads as "no responses"."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta {target}"
        )

        response = await client.get(
            _responses_url(survey["id"], "export"),
            params={"establishment_id": str(db_establishment.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_field_worker_forbidden(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)

        response = await client.get(
            _responses_url(survey["id"], "export"),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_unknown_survey_returns_404(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        unknown_id = uuid.uuid4()
        response = await client.get(
            _responses_url(unknown_id, "export"),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_export_with_no_responses_still_succeeds(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)

        response = await client.get(
            _responses_url(survey["id"], "export"),
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK
