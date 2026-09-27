import uuid
from io import BytesIO

import pytest
from fastapi import status
from openpyxl import load_workbook

from tests.integration.helpers import responses_url
from tests.integration.surveys.responses.helpers import (
    _create_survey,
    _response_body,
    _valid_answers,
)


class TestExport:
    async def test_admin_can_export_xlsx(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        await client.post(
            responses_url(survey["id"]),
            json=_response_body(survey["id"], _valid_answers(survey)),
            headers=headers,
        )

        response = await client.get(
            responses_url(survey["id"], "export"),
            headers=headers,
        )

        assert response.status_code == status.HTTP_200_OK
        assert "spreadsheetml" in response.headers["content-type"]
        assert response.content[:2] == b"PK"  # xlsx is a zip archive

    async def test_response_number_range_export_includes_duplicates(
        self, client, db_establishment, admin_user
    ):
        """The number range export retains all matching duplicate labels."""
        headers, _ = admin_user
        survey = await _create_survey(client, headers)
        for response_number in (74, 74, 75):
            response = await client.post(
                responses_url(survey["id"]),
                json=_response_body(
                    survey["id"],
                    _valid_answers(survey),
                    response_number=response_number,
                ),
                headers=headers,
            )
            assert response.status_code == status.HTTP_201_CREATED

        response = await client.get(
            responses_url(survey["id"], "export"),
            params={"response_number_from": 74, "response_number_to": 74},
            headers=headers,
        )

        assert response.status_code == status.HTTP_200_OK
        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook["Respuestas"].iter_rows(min_row=2, values_only=True))
        assert [row[0] for row in rows] == [74, 74]

    async def test_filters_by_establishment(self, client, db_establishment, admin_user):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="establishment")
        await client.post(
            responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                establishment_id=str(db_establishment.id),
            ),
            headers=headers,
        )

        response = await client.get(
            responses_url(survey["id"], "export"),
            params={"establishment_id": str(db_establishment.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK

    async def test_filters_by_subject(self, client, admin_user, db_subject):
        headers, _ = admin_user
        survey = await _create_survey(client, headers, response_target="subject")
        await client.post(
            responses_url(survey["id"]),
            json=_response_body(
                survey["id"],
                _valid_answers(survey),
                subject_id=str(db_subject.id),
            ),
            headers=headers,
        )

        response = await client.get(
            responses_url(survey["id"], "export"),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK

    @pytest.mark.parametrize("target", ["establishment", "anonymous"])
    async def test_subject_filter_on_non_subject_survey_returns_400(
        self, client, admin_user, db_subject, target
    ):
        """Nothing in this survey carries a subject, so the filter would export
        an empty workbook that reads as "no responses"."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta sujeto {target}"
        )

        response = await client.get(
            responses_url(survey["id"], "export"),
            params={"subject_id": str(db_subject.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    @pytest.mark.parametrize("target", ["subject", "anonymous"])
    async def test_establishment_filter_on_non_establishment_survey_returns_400(
        self, client, db_establishment, admin_user, target
    ):
        """Nothing in this survey carries an establishment, so the filter would
        export an empty workbook that reads as "no responses"."""
        headers, _ = admin_user
        survey = await _create_survey(
            client, headers, response_target=target, name=f"Encuesta {target}"
        )

        response = await client.get(
            responses_url(survey["id"], "export"),
            params={"establishment_id": str(db_establishment.id)},
            headers=headers,
        )
        assert response.status_code == status.HTTP_400_BAD_REQUEST

    async def test_field_worker_forbidden(
        self, client, db_establishment, admin_user, field_worker_with_access
    ):
        admin_headers, _ = admin_user
        fw_headers, _ = field_worker_with_access
        survey = await _create_survey(client, admin_headers)

        response = await client.get(
            responses_url(survey["id"], "export"),
            headers=fw_headers,
        )
        assert response.status_code == status.HTTP_403_FORBIDDEN

    async def test_unknown_survey_returns_404(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        unknown_id = uuid.uuid4()
        response = await client.get(
            responses_url(unknown_id, "export"),
            headers=headers,
        )
        assert response.status_code == status.HTTP_404_NOT_FOUND

    async def test_export_with_no_responses_still_succeeds(
        self, client, db_establishment, admin_user
    ):
        headers, _ = admin_user
        survey = await _create_survey(client, headers)

        response = await client.get(
            responses_url(survey["id"], "export"),
            headers=headers,
        )
        assert response.status_code == status.HTTP_200_OK
